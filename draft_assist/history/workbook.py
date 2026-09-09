"""The export: two sheets, at the user's request.

ONE tab is the raw table — one row per match, under an autofilter, which
is the sheet you sort and pivot yourself. The OTHER is the whole analysis:
the run's metadata, every written finding, then each block's buckets with
its counts, deviations and sigmas stacked under a heading. The browser
version wrote one sheet per analysis and a summary; thirteen tabs is a
worse way to read the same thing than one you can scroll.

Every bucket is exported, INCLUDING the single-game ones the on-screen
table hides — the screen is protecting you from reading noise, the
workbook is the record.
"""

from datetime import datetime

from .analyse import MIN_BUCKET, MIN_DISPLAY, SIGMA_CAT
from .shape import GAME_MODE, LANE_ROLE, LOBBY_TYPE, WEEKDAY

RAW_HEADERS = ["Match ID", "Date", "Time", "Hour", "Weekday", "Hero", "Result",
               "Side", "Duration (min)", "Kills", "Deaths", "Assists", "GPM",
               "XPM", "Last hits", "Denies", "Hero damage", "Tower damage",
               "Level", "Party size", "Lobby", "Mode", "Lane role", "Session",
               "Position in session", "Previous result"]


def raw_rows(report) -> list:
    out = [list(RAW_HEADERS)]
    for match in report.matches:
        when = match.when
        out.append([
            match.match_id, when.strftime("%Y-%m-%d"), when.strftime("%H:%M"),
            when.hour, WEEKDAY[when.weekday()], match.hero,
            "Win" if match.win else "Loss",
            "Radiant" if match.radiant else "Dire", round(match.minutes, 2),
            match.kills, match.deaths, match.assists, match.gold_per_min,
            match.xp_per_min, match.last_hits, match.denies,
            match.hero_damage, match.tower_damage, match.level,
            "unknown" if match.party_size is None else match.party_size,
            LOBBY_TYPE.get(match.lobby_type,
                           "unknown" if match.lobby_type is None
                           else f"Lobby {match.lobby_type}"),
            GAME_MODE.get(match.game_mode,
                          "unknown" if match.game_mode is None
                          else f"Mode {match.game_mode}"),
            LANE_ROLE.get(match.lane_role, "unparsed"),
            match.session, match.position, match.previous])
    return out


def analysis_rows(report) -> list:
    """Metadata, findings, then every block's buckets, stacked."""
    options = report.options
    period = report.period
    out = [
        ["Dota deviation report"], [],
        ["Account ID", options.account_id],
        ["Read as", report.how or ""],
        ["Name", report.name or ""],
        ["Generated", report.ran_at.strftime("%Y-%m-%d %H:%M:%S")
         + " local machine time"],
        ["Requested window", options.window_label],
        ["Requested cap", options.cap],
        ["Exclude Turbo", "yes" if options.no_turbo else "no"],
        ["Ranked only", "yes" if options.ranked_only else "no"], [],
        ["Matches returned by API", report.returned],
        ["Matches after filters", report.n],
        ["Wins", report.wins], ["Losses", report.n - report.wins],
        ["Baseline win rate", round(report.baseline, 5)],
        ["Period covered", f"{period[0]} to {period[1]}"],
        ["Play sessions", report.sessions], [],
        ["Dropped, under 5 minutes", report.dropped.get("short", 0)],
        ["Dropped, outside window", report.dropped.get("window", 0)],
        ["Dropped, Turbo", report.dropped.get("turbo", 0)],
        ["Dropped, not ranked", report.dropped.get("unranked", 0)],
        ["Dropped, malformed", report.dropped.get("malformed", 0)], [],
        ["Significance floor (sigma)", SIGMA_CAT],
        ["Minimum games to appear in a table", MIN_DISPLAY],
        ["Minimum games for a finding", MIN_BUCKET],
        ["Note", "Twelve analyses run at once, so some buckets clear the bar "
                 "by chance alone. A finding is a hypothesis to test against "
                 "the next hundred games, not a conclusion."], [],
        ["FINDINGS"], ["Analysis", "Sigma", "Finding"]]

    for block in report.blocks:
        if not block.findings:
            out.append([block.name, "",
                        "No split clears the significance floor. Noise."])
            continue
        for finding in block.findings:
            out.append([block.name, round(finding.sigma, 3), finding.text])

    for block in report.blocks:
        out += [[], [], [block.name.upper()], [block.desc]]
        if block.kind == "items":
            out.append(["Hero", "Hero games", "Games with item data",
                        "Hero win rate", "Item", "Games with item", "Wins",
                        "Win rate", "Deviation", "Standard error", "Sigma",
                        "Eligible"])
            for group in block.groups:
                if group.baseline is None:
                    out.append([group.hero, group.games, 0])
                    continue
                for row in group.rows:
                    out.append([group.hero, group.games, group.measured,
                                round(group.baseline, 5), row.key, row.n,
                                row.wins, round(row.rate, 5),
                                round(row.delta, 5), round(row.se, 5),
                                round(row.sigma, 3),
                                "yes" if row.eligible else "no"])
        elif block.kind == "metric":
            out.append(["Hero", "Games", f"Mean {block.unit}",
                        "Deviation from overall", "Standard error", "Sigma",
                        "Eligible for finding", "Shown in table"])
            out.append(["ALL (your overall)", block.covered,
                        "" if block.datum is None else round(block.datum, 3),
                        0])
            out.append(["Matches carrying this field",
                        f"{block.covered} of {block.total}"])
            for row in block.rows:
                out.append([row.key, row.n, round(row.mean, 3),
                            round(row.delta, 3), round(row.se, 3),
                            round(row.sigma, 3),
                            "yes" if row.eligible else "no",
                            "yes" if row.n >= MIN_DISPLAY else "no"])
        else:
            out.append(["Bucket", "Games", "Wins", "Losses", "Win rate",
                        "Deviation from datum", "Standard error", "Sigma",
                        "Eligible for finding", "Shown in table"])
            for row in block.rows:
                out.append([row.key, row.n, row.wins, row.n - row.wins,
                            round(row.rate, 5), round(row.delta, 5),
                            round(row.se, 5), round(row.sigma, 3),
                            "yes" if (row.eligible
                                      and row.key not in block.no_finding)
                            else "no",
                            "yes" if row.n >= MIN_DISPLAY else "no"])
    return out


def filename(report) -> str:
    stamp = report.ran_at.strftime("%Y-%m-%d")
    return f"dota-deviation-{report.options.account_id}-{stamp}.xlsx"


def write(report, path) -> str:
    """Write the workbook. Returns the path it went to.

    openpyxl rather than a CSV pair: the raw sheet is meant to be sorted
    and filtered, and an autofilter is the one thing a CSV cannot carry.
    """
    from openpyxl import Workbook
    from openpyxl.utils import get_column_letter

    book = Workbook()
    raw = book.active
    raw.title = "Raw matches"
    rows = raw_rows(report)
    for row in rows:
        raw.append(row)
    if len(rows) > 1:
        raw.auto_filter.ref = (
            f"A1:{get_column_letter(len(RAW_HEADERS))}{len(rows)}")
    raw.freeze_panes = "A2"

    sheet = book.create_sheet("Analysis")
    for row in analysis_rows(report):
        sheet.append(row)
    # Wide enough to read the finding sentences without dragging columns.
    for column, width in (("A", 34), ("B", 14), ("C", 92)):
        sheet.column_dimensions[column].width = width

    book.save(str(path))
    return str(path)
